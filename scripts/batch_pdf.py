"""Batch-создание PDF отчётов из Валидных final Excel-отчётов.

Resumable:
  - final Excel -> PDF (reports/final_transcript_report_<DATE>.pdf)
  - Если PDF уже существует и валиден (открывается, >0 страниц) - SKIP.
  - Использует existing PDF logic (offline_analysis.pdf_report.generate_pdf_report).
  - PDF строится ТОЛЬКО ИЗ FINAL REPORT (не из raw).

НЕ дублирует логику формирования отчёта — пересказывает строки final Excel
в `rows` для `generate_pdf_report`.
"""
from __future__ import annotations

import datetime as _dt
import sys
from pathlib import Path

import openpyxl

# Импорт PDF-генератора из проекта (один и тот же формат, та же логика).
_REPO_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(_REPO_ROOT))
from offline_analysis.pdf_report import generate_pdf_report  # noqa: E402

from pypdf import PdfReader  # noqa: E402

REPORTS_DIR = _REPO_ROOT / "reports"


def _read_final_rows(xlsx: Path) -> list[dict]:
    """Читает sheet `report` -> list[dict] (ключи = колонки)."""
    wb = openpyxl.load_workbook(xlsx, read_only=True)
    try:
        if "report" not in wb.sheetnames:
            raise RuntimeError(f"no 'report' sheet in {xlsx}")
        ws = wb["report"]
        it = ws.iter_rows(values_only=True)
        header = next(it)
        header = [str(h) if h is not None else "" for h in header]
        rows: list[dict] = []
        for raw in it:
            if raw is None or all(c is None for c in raw):
                continue
            rows.append({h: raw[i] for i, h in enumerate(header) if h})
    finally:
        wb.close()
    return rows


def _pdf_is_valid(pdf: Path) -> bool:
    if not pdf.exists() or pdf.stat().st_size <= 0:
        return False
    try:
        r = PdfReader(str(pdf))
        return len(r.pages) > 0
    except Exception:  # noqa: BLE001
        return False


def main() -> int:
    dates = sorted({
        p.name.replace("final_transcript_report_", "").replace(".xlsx", "")
        for p in REPORTS_DIR.glob("final_transcript_report_*.xlsx")
    })
    n_ok = n_skip = n_fail = n_empty = 0
    for d in dates:
        xlsx = REPORTS_DIR / f"final_transcript_report_{d}.xlsx"
        pdf = REPORTS_DIR / f"final_transcript_report_{d}.pdf"

        # 1) Final Excel отсутствует/пустой -> не делаем PDF
        rows = _read_final_rows(xlsx)
        if not rows:
            n_empty += 1
            print(f"[{d}] SKIP (final has 0 rows)")
            continue

        # 2) PDF уже валиден -> SKIP
        if _pdf_is_valid(pdf):
            n_skip += 1
            print(f"[{d}] SKIP (pdf already valid)")
            continue

        # 3) Создаём
        try:
            out = generate_pdf_report(rows, str(pdf), report_date=d)
            n_ok += 1
            print(f"[{d}] OK -> {out}")
        except Exception as exc:  # noqa: BLE001
            n_fail += 1
            print(f"[{d}] FAIL {exc!r}", file=sys.stderr)

    print()
    print(f"done: ok={n_ok} skip={n_skip} empty={n_empty} fail={n_fail}")
    return 0 if n_fail == 0 else 1


if __name__ == "__main__":
    sys.exit(main())
