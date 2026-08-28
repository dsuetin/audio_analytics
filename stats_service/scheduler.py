from __future__ import annotations

import argparse
import logging
import os
import subprocess
import sys
import time as time_module
from datetime import date, datetime, time, timedelta
from pathlib import Path
from zoneinfo import ZoneInfo


logger = logging.getLogger(__name__)


def env_value(name: str, default: str) -> str:
    value = os.getenv(name, "").strip()
    return value or default


def resolve_report_date(zone: ZoneInfo) -> date:
    return datetime.now(zone).date() - timedelta(days=1)


def build_output_path(output_dir: Path, report_date: date) -> Path:
    return output_dir / f"transcript_report_{report_date.isoformat()}.xlsx"


def seconds_until_next_midnight(zone: ZoneInfo) -> float:
    now = datetime.now(zone)
    next_midnight = datetime.combine(now.date() + timedelta(days=1), time.min, tzinfo=zone)
    return max((next_midnight - now).total_seconds(), 0.0)


def run_daily_report(
    report_date: date,
    timezone_name: str,
    output_dir: Path,
    postgres_host: str,
    postgres_port: str,
    postgres_user: str,
    postgres_password: str,
    postgres_db: str,
    offline_enabled: bool | None = None,
) -> None:
    repo_root = Path(__file__).resolve().parents[1]
    script_path = repo_root / "stats_service" / "export_daily_transcript_report.py"
    output_path = build_output_path(output_dir, report_date)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    command = [
        sys.executable,
        str(script_path),
        "--date",
        report_date.isoformat(),
        "--timezone",
        timezone_name,
        "--output",
        str(output_path),
    ]

    logger.info("Running report for %s", report_date.isoformat())
    child_env = os.environ.copy()
    child_env.update(
        {
            "POSTGRES_HOST": postgres_host,
            "POSTGRES_PORT": postgres_port,
            "POSTGRES_USER": postgres_user,
            "POSTGRES_PASSWORD": postgres_password,
            "POSTGRES_DB": postgres_db,
        }
    )
    result = subprocess.run(
        command,
        cwd=str(repo_root),
        check=False,
        capture_output=True,
        text=True,
        env=child_env,
    )

    if result.stdout:
        logger.info("report stdout:\n%s", result.stdout.rstrip())
    if result.stderr:
        logger.warning("report stderr:\n%s", result.stderr.rstrip())

    if result.returncode != 0:
        # RAW generation failed: do NOT touch offline analysis, keep the exit non-zero.
        # The (possibly stale/absent) raw file is left in place for a manual re-run.
        raise RuntimeError(
            f"RAW REPORT: FAILED | generation failed with exit code {result.returncode} | {output_path}"
        )

    # RAW generation exited 0 — but we do not trust the code alone: the file must exist.
    ensure_raw_report_created(output_path)

    logger.info("RAW REPORT: OK | Report generated: %s and %s", output_path, output_path.with_suffix(".pdf"))

    # [2/4]..[4/4] Sequentially chain the existing offline analysis on the SAME raw file.
    # The scheduler already waits synchronously; we keep it sequential (no background / no new cron).
    run_offline_analysis(
        repo_root=repo_root,
        report_date=report_date,
        raw_path=output_path,
        output_dir=output_dir,
        enabled=offline_enabled,
    )


def ensure_raw_report_created(output_path: Path) -> None:
    """Required rule: 'exit 0 but no file' is still an error and offline must NOT run.

    Raises RuntimeError if the raw report file is absent (or is an empty zero-byte file).
    """
    if not output_path.exists():
        raise RuntimeError(
            f"RAW REPORT: FAILED | exit code was 0 but the raw report was not created: {output_path}"
        )
    try:
        if output_path.stat().st_size == 0:
            raise RuntimeError(
                f"RAW REPORT: FAILED | file was created but is empty (0 bytes): {output_path}"
            )
    except OSError as exc:
        raise RuntimeError(f"RAW REPORT: FAILED | cannot stat raw report {output_path}: {exc}") from exc


def _truthy(name: str, default: bool = False) -> bool:
    value = os.getenv(name, "").strip().lower()
    if value == "":
        return default
    return value in ("1", "true", "yes", "y", "on")


def run_offline_analysis(
    repo_root: Path,
    report_date: date,
    raw_path: Path,
    output_dir: Path,
    enabled: bool | None = None,
) -> None:
    """Запускает существующий offline-анализ (offline_analysis/run.py) по-очерёдно,
    после успешной генерации raw-отчёта. Ждёт завершения синхронно.

    Ошибки оффлайна НЕ удаляют raw-отчёт: raw остаётся для ручного повторного запуска.
    """
    # --- default: ON unless the existing knob says otherwise -------------------
    if enabled is None:
        enabled = _truthy("OFFLINE_ANALYSIS_ENABLED", True)
    if not enabled:
        logger.info("OFFLINE ANALYSIS: SKIPPED (OFFLINE_ANALYSIS_ENABLED=false)")
        return

    force = _truthy("FORCE_OFFLINE", False)
    final_path = output_dir / f"final_transcript_report_{report_date.isoformat()}.xlsx"

    # --- idempotency (existing style: env knob, not a new config system) ------
    if not force and final_path.exists():
        try:
            problems = validate_final_report(final_path)
        except Exception as exc:  # noqa: BLE001 - corrupted file => treat as "not valid"
            logger.warning("existing final report is not valid, re-running offline: %s", exc)
            problems = [f"unreadable: {exc}"]
        if not problems:
            logger.info("OFFLINE ANALYSIS: SKIPPED (final report already valid: %s)", final_path)
            return

    offline_script = repo_root / "offline_analysis" / "run.py"
    if not offline_script.exists():
        raise RuntimeError(
            f"OFFLINE ANALYSIS: FAILED | offline analyzer not found in this deployment: {offline_script}"
        )

    command = [
        sys.executable,
        str(offline_script),
        "--input",
        str(raw_path),
        "--output",
        str(final_path),
        # debug path is chosen by run.py from the output stem: final_..._debug.json
    ]

    # Optional parallelism for the LLM step (existing env-knob style; unset => safe default in run.py).
    workers = os.getenv("OFFLINE_WORKERS", "").strip()
    if workers:
        command += ["--workers", workers]

    logger.info("OFFLINE: starting analyzer for %s (final=%s, workers=%s)", raw_path, final_path, workers or "default")
    # Reuse the existing env-based configuration (model/host/timeout) — see offline_analysis/llm.py.
    child_env = os.environ.copy()
    child_env.setdefault("OFFLINE_LLM_MODEL", "qwen3.8:27b")
    # OLLAMA_HOST may already be configured; run.py/llm.py read it from the environment.

    # We do not capture and hide the output: surface it as the daily pipeline log.
    result = subprocess.run(
        command,
        cwd=str(repo_root),
        check=False,
        text=True,
        env=child_env,
    )

    if result.returncode != 0:
        # RAW report is intentionally left in place for a manual re-run.
        raise RuntimeError(
            f"OFFLINE ANALYSIS: FAILED | exit code {result.returncode} | raw kept: {raw_path} | final: {final_path}"
        )

    # exit 0 is not enough — the final file must physically exist.
    if not final_path.exists():
        raise RuntimeError(
            f"OFFLINE ANALYSIS: FAILED | offline reported success but the final report was not created: {final_path}"
        )

    # [4/4] Validate that the final report is actually readable and well-formed.
    problems = validate_final_report(final_path)
    if problems:
        raise RuntimeError(
            "OFFLINE ANALYSIS: FAILED | final report is invalid: " + "; ".join(problems)
        )

    logger.info("OFFLINE ANALYSIS: OK | Final report created and valid: %s", final_path)
    logger.info("PIPELINE: DONE for %s", report_date.isoformat())


def validate_final_report(final_path: Path) -> list[str]:
    """Минимальная проверка final Excel: лист report, ключевые колонки.
    Возвращает список проблем (пустой список = OK)."""
    if not final_path.exists():
        return [f"file not found: {final_path}"]
    required_columns = {
        "store_id",
        "client_id",
        "dialog_type",
        "model_reasoning",
        "is_sale",
    }
    try:
        from openpyxl import load_workbook

        wb = load_workbook(final_path, read_only=True)
        try:
            if "report" not in wb.sheetnames:
                return [f"missing sheet 'report' (sheets={wb.sheetnames})"]
            ws = wb["report"]
            header = [ (c.value) for c in next(ws.iter_rows(min_row=1, max_row=1)) ]
            missing = required_columns - set(header)
            if missing:
                return [f"missing columns: {sorted(missing)}"]
            return []
        finally:
            wb.close()
    except Exception as exc:  # noqa: BLE001
        return [f"failed to open/validate: {exc}"]


def main() -> None:
    parser = argparse.ArgumentParser(description="Run the daily statistics export at midnight.")
    parser.add_argument(
        "--timezone",
        default=env_value("REPORT_TIMEZONE", "Europe/Moscow"),
        help="Timezone used to calculate midnight and the previous day.",
    )
    parser.add_argument(
        "--date",
        help="Optional report date in YYYY-MM-DD format. Defaults to yesterday in the chosen timezone.",
    )
    parser.add_argument(
        "--output-dir",
        default=env_value("REPORT_OUTPUT_DIR", "reports"),
        help="Directory for generated report files.",
    )
    parser.add_argument(
        "--postgres-host",
        default=env_value("POSTGRES_HOST", "localhost"),
        help="PostgreSQL host for local runs. Use postgres inside Docker.",
    )
    parser.add_argument(
        "--postgres-port",
        default=env_value("POSTGRES_PORT", "5432"),
        help="PostgreSQL port.",
    )
    parser.add_argument(
        "--postgres-user",
        default=env_value("POSTGRES_USER", "speech"),
        help="PostgreSQL user.",
    )
    parser.add_argument(
        "--postgres-password",
        default=env_value("POSTGRES_PASSWORD", "speech"),
        help="PostgreSQL password.",
    )
    parser.add_argument(
        "--postgres-db",
        default=env_value("POSTGRES_DB", "speech_db"),
        help="PostgreSQL database name.",
    )
    parser.add_argument(
        "--run-on-startup",
        action="store_true",
        help="Generate the previous day report immediately, then continue with midnight scheduling.",
    )
    parser.add_argument(
        "--once",
        action="store_true",
        help="Generate one report and exit instead of waiting for midnight.",
    )
    args = parser.parse_args()

    logging.basicConfig(
        level=os.getenv("LOG_LEVEL", "INFO").upper(),
        format="%(asctime)s %(levelname)s %(name)s: %(message)s",
    )

    zone = ZoneInfo(args.timezone)
    output_dir = Path(args.output_dir)
    report_date = datetime.strptime(args.date, "%Y-%m-%d").date() if args.date else resolve_report_date(zone)

    if args.run_on_startup or args.once:
        run_daily_report(
            report_date,
            args.timezone,
            output_dir,
            args.postgres_host,
            args.postgres_port,
            args.postgres_user,
            args.postgres_password,
            args.postgres_db,
        )
        if args.once:
            return

    while True:
        sleep_seconds = seconds_until_next_midnight(zone)
        logger.info("Sleeping %.1f seconds until midnight", sleep_seconds)
        time_module.sleep(sleep_seconds)

        report_date = resolve_report_date(zone)
        try:
            run_daily_report(
                report_date,
                args.timezone,
                output_dir,
                args.postgres_host,
                args.postgres_port,
                args.postgres_user,
                args.postgres_password,
                args.postgres_db,
            )
        except Exception:
            logger.exception("Daily report failed")
            time_module.sleep(60)


if __name__ == "__main__":
    main()
