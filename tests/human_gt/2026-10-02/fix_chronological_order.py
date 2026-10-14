#!/usr/bin/env python3
"""
Исправление хронологического порядка в TIMELINE_AI_SIMPLE.

Проблема: GT время (HH:MM:SS) и AI время (YYYY-MM-DD HH:MM) сортируются неправильно.
Решение: Нормализовать все времена к формату минут от начала дня и сортировать по числам.
"""

from __future__ import annotations

import openpyxl
from datetime import datetime
from pathlib import Path


def parse_time_to_minutes(time_str: str | None) -> float | None:
    """
    Convert various time formats to minutes from start of day for sorting.
    
    Supported formats:
    - '2026-10-02 08:38' (AI format)
    - '11:10:35' (GT format)
    - '08:38:28' (full time)
    """
    if not time_str:
        return None
    
    time_str = str(time_str).strip()
    
    try:
        # Try AI format: '2026-10-02 08:38'
        if '-' in time_str:
            # Extract time part after date
            time_part = time_str.split(' ')[1] if ' ' in time_str else time_str
            parts = time_part.split(':')
            if len(parts) >= 2:
                hours = int(parts[0])
                minutes = int(parts[1])
                seconds = int(parts[2]) if len(parts) > 2 else 0
                return hours * 60 + minutes + seconds / 60
        
        # Try GT format: '11:10:35' or '08:38'
        parts = time_str.split(':')
        if len(parts) >= 2:
            hours = int(parts[0])
            minutes = int(parts[1])
            seconds = int(parts[2]) if len(parts) > 2 else 0
            return hours * 60 + minutes + seconds / 60
        
        return None
    except (ValueError, IndexError):
        return None


def fix_chronological_order():
    """Fix chronological order in the simple timeline report."""
    
    source_file = Path("tests/human_gt/2026-10-02/TIMELINE_AI_SIMPLE_2026-10-02.xlsx")
    output_file = Path("tests/human_gt/2026-10-02/TIMELINE_AI_SIMPLE_SORTED_2026-10-02.xlsx")
    
    # Load source workbook
    wb_source = openpyxl.load_workbook(source_file)
    
    # Create new workbook
    wb_new = openpyxl.Workbook()
    default_sheet = wb_new.active
    wb_new.remove(default_sheet)
    
    print("ИСПРАВЛЕНИЕ ХРОНОЛОГИЧЕСКОГО ПОРЯДКА")
    print("=" * 80)
    
    stores_processed = 0
    total_rows_sorted = 0
    rows_with_time_issues = 0
    
    for sheet_name in wb_source.sheetnames:
        if sheet_name == "Сводка":
            # Copy summary sheet as-is
            ws_source = wb_source[sheet_name]
            ws_new = wb_new.create_sheet(sheet_name)
            
            for row in ws_source.iter_rows():
                for cell in row:
                    if cell.value is not None:
                        ws_new.cell(row=cell.row, column=cell.column, value=cell.value)
            continue
        
        stores_processed += 1
        ws_source = wb_source[sheet_name]
        
        # Collect all data rows
        rows_data = []
        for row in range(2, ws_source.max_row + 1):
            time_val = ws_source.cell(row=row, column=1).value
            ai_mission = ws_source.cell(row=row, column=2).value
            ai_outcome = ws_source.cell(row=row, column=3).value
            comment = ws_source.cell(row=row, column=4).value
            
            # Skip empty rows
            if not time_val:
                continue
            
            # Parse time for sorting
            time_minutes = parse_time_to_minutes(time_val)
            
            if time_minutes is None:
                rows_with_time_issues += 1
                # Put rows with unknown time at the end
                time_minutes = 999999  # Sort to end
            
            rows_data.append({
                'time_minutes': time_minutes,
                'time_val': time_val,
                'ai_mission': ai_mission,
                'ai_outcome': ai_outcome,
                'comment': comment,
                'has_time_issue': time_minutes == 999999
            })
        
        # Sort by time
        rows_data.sort(key=lambda x: x['time_minutes'])
        
        # Create new sheet
        ws_new = wb_new.create_sheet(sheet_name)
        
        # Write header
        headers = ["Время", "AI — миссия", "AI — результат", "Комментарий"]
        for col_idx, header in enumerate(headers, 1):
            cell = ws_new.cell(row=1, column=col_idx, value=header)
            cell.font = openpyxl.styles.Font(bold=True)
            cell.fill = openpyxl.styles.PatternFill(start_color="DDDDDD", fill_type="solid")
        
        # Write sorted data
        for row_idx, row_data in enumerate(rows_data, 2):
            ws_new.cell(row=row_idx, column=1, value=row_data['time_val'])
            ws_new.cell(row=row_idx, column=2, value=row_data['ai_mission'])
            ws_new.cell(row=row_idx, column=3, value=row_data['ai_outcome'])
            ws_new.cell(row=row_idx, column=4, value=row_data['comment'])
            
            # Highlight rows with time issues
            if row_data['has_time_issue']:
                ws_new.cell(row=row_idx, column=1).fill = openpyxl.styles.PatternFill(
                    start_color="FFCCCC", fill_type="solid"
                )
                # Add note to comment
                if row_data['comment']:
                    row_data['comment'] += " [время не распознано]"
                else:
                    row_data['comment'] = "время не распознано"
                ws_new.cell(row=row_idx, column=4, value=row_data['comment'])
        
        # Set column widths
        ws_new.column_dimensions["A"].width = 18
        ws_new.column_dimensions["B"].width = 25
        ws_new.column_dimensions["C"].width = 20
        ws_new.column_dimensions["D"].width = 40
        
        # Wrap text for comment column
        for row_idx in range(2, len(rows_data) + 2):
            ws_new.cell(row=row_idx, column=4).alignment = openpyxl.styles.Alignment(wrap_text=True)
        
        total_rows_sorted += len(rows_data)
        
        print(f"  {sheet_name}: {len(rows_data)} записей отсортировано")
    
    # Update summary sheet
    if "Сводка" in wb_new.sheetnames:
        ws_summary = wb_new["Сводка"]
        ws_summary.cell(row=4, column=1, value="Дата сортировки:")
        ws_summary.cell(row=4, column=2, value=datetime.now().strftime("%Y-%m-%d %H:%M:%S"))
        ws_summary.cell(row=6, column=1, value="Записей с проблемами времени:")
        ws_summary.cell(row=6, column=2, value=rows_with_time_issues)
    
    # Save workbook
    wb_new.save(output_file)
    
    print(f"\n{'=' * 80}")
    print(f"Файл сохранён: {output_file.absolute()}")
    print(f"Обработано магазинов: {stores_processed}")
    print(f"Всего записей отсортировано: {total_rows_sorted}")
    print(f"Записей с проблемами времени: {rows_with_time_issues}")
    print(f"\nПричина неправильной сортировки:")
    print(f"  GT время: формат 'HH:MM:SS' (например, '11:10:35')")
    print(f"  AI время: формат 'YYYY-MM-DD HH:MM' (например, '2026-10-02 08:38')")
    print(f"  При текстовой сортировке '11:10:35' < '2026-10-02' (символ '1' < '2')")
    print(f"  Решение: конвертация в минуты от начала дня и числовая сортировка")


if __name__ == "__main__":
    fix_chronological_order()
