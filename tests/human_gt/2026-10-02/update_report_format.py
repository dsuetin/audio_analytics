#!/usr/bin/env python3
"""
Доработка отчёта AI vs GT:
1. Единый формат времени ЧЧ:ММ:СС
2. Замена "1. + 2. +" на жирную галочку ✔
3. Добавление итогов внизу каждого листа
"""

from __future__ import annotations

import openpyxl
from openpyxl.styles import Font, Alignment
from datetime import datetime
from pathlib import Path


def parse_time_to_minutes(time_str: str) -> tuple[float, str]:
    """
    Parse time string and return (minutes_from_start, formatted_time).
    
    Supported formats:
    - '2026-10-02 08:38' (AI format)
    - '11:10:35' (GT format)
    """
    time_str = str(time_str).strip()
    
    # Extract time part
    if '-' in time_str:
        # AI format: '2026-10-02 08:38'
        time_part = time_str.split(' ')[1] if ' ' in time_str else time_str
    else:
        # GT format: '11:10:35'
        time_part = time_str
    
    # Parse hours, minutes, seconds
    parts = time_part.split(':')
    hours = int(parts[0])
    minutes = int(parts[1])
    seconds = int(parts[2]) if len(parts) > 2 else 0
    
    # Calculate minutes from start of day
    time_minutes = hours * 60 + minutes + seconds / 60
    
    # Format as HH:MM:SS
    formatted_time = f"{hours:02d}:{minutes:02d}:{seconds:02d}"
    
    return time_minutes, formatted_time


def is_positive_confirmation(comment: str) -> bool:
    """Check if comment is a positive confirmation (1.+ 2.+)"""
    if not comment:
        return False
    
    comment_lower = str(comment).lower().strip()
    
    # Check for various forms of positive confirmation
    patterns = [
        "1. + 2. +",
        "1.+ 2.+",
        "1.+ 2. +",
        "1. + 2.+",
        "1+2+",
    ]
    
    return any(pattern in comment_lower for pattern in patterns)


def create_updated_report():
    """Create updated report with unified format, checkmarks, and totals."""
    
    source_file = Path("tests/human_gt/2026-10-02/TIMELINE_AI_SIMPLE_SORTED_2026-10-02.xlsx")
    output_file = Path("tests/human_gt/2026-10-02/TIMELINE_AI_SIMPLE_FINAL_2026-10-02.xlsx")
    
    # Load source workbook
    wb_source = openpyxl.load_workbook(source_file)
    
    # Create new workbook
    wb_new = openpyxl.Workbook()
    default_sheet = wb_new.active
    wb_new.remove(default_sheet)
    
    print("СОЗДАНИЕ ОБНОВЛЁННОГО ОТЧЁТА")
    print("=" * 80)
    
    stores_processed = 0
    total_rows = 0
    total_checkmarks = 0
    
    for sheet_name in wb_source.sheetnames:
        if sheet_name == "Сводка":
            # Copy summary sheet with updates
            ws_source = wb_source[sheet_name]
            ws_new = wb_new.create_sheet(sheet_name)
            
            for row in ws_source.iter_rows():
                for cell in row:
                    if cell.value is not None:
                        ws_new.cell(row=cell.row, column=cell.column, value=cell.value)
            
            # Add update info
            ws_new.cell(row=ws_new.max_row + 2, column=1, value="Обновление:")
            ws_new.cell(row=ws_new.max_row + 3, column=1, value="Формат времени: ЧЧ:ММ:СС")
            ws_new.cell(row=ws_new.max_row + 4, column=1, value="Галочки: ✔ для '1.+ 2.+'")
            ws_new.cell(row=ws_new.max_row + 5, column=1, value="Итоги: добавлены внизу каждого листа")
            ws_new.cell(row=ws_new.max_row + 6, column=1, value="Дата обновления:")
            ws_new.cell(row=ws_new.max_row + 7, column=2, value=datetime.now().strftime("%Y-%m-%d %H:%M:%S"))
            
            ws_new.column_dimensions["A"].width = 30
            ws_new.column_dimensions["B"].width = 50
            continue
        
        stores_processed += 1
        ws_source = wb_source[sheet_name]
        
        # Create new sheet
        ws_new = wb_new.create_sheet(sheet_name)
        
        # Write header
        headers = ["Время", "AI — миссия", "AI — результат", "Комментарий"]
        for col_idx, header in enumerate(headers, 1):
            cell = ws_new.cell(row=1, column=col_idx, value=header)
            cell.font = Font(bold=True)
            cell.fill = openpyxl.styles.PatternFill(start_color="DDDDDD", fill_type="solid")
        
        # Process data rows
        store_rows = 0
        store_checkmarks = 0
        
        for row in range(2, ws_source.max_row + 1):
            time_val = ws_source.cell(row=row, column=1).value
            ai_mission = ws_source.cell(row=row, column=2).value
            ai_outcome = ws_source.cell(row=row, column=3).value
            comment = ws_source.cell(row=row, column=4).value
            
            # Skip empty rows
            if not time_val:
                continue
            
            store_rows += 1
            total_rows += 1
            
            # Convert time to unified format
            try:
                _, formatted_time = parse_time_to_minutes(time_val)
            except (ValueError, IndexError):
                formatted_time = str(time_val)  # Keep original if parsing fails
            
            # Check for positive confirmation and replace with checkmark
            if is_positive_confirmation(comment):
                comment = "✔"
                store_checkmarks += 1
                total_checkmarks += 1
            
            # Write to new sheet
            ws_new.cell(row=store_rows + 1, column=1, value=formatted_time)
            ws_new.cell(row=store_rows + 1, column=2, value=ai_mission)
            ws_new.cell(row=store_rows + 1, column=3, value=ai_outcome)
            
            # Write comment with special formatting for checkmark
            comment_cell = ws_new.cell(row=store_rows + 1, column=4, value=comment)
            if comment == "✔":
                comment_cell.font = Font(bold=True, size=14)
                comment_cell.alignment = Alignment(horizontal="center")
            else:
                comment_cell.alignment = Alignment(wrap_text=True)
        
        # Add totals at bottom
        totals_row = store_rows + 3
        totals_cell = ws_new.cell(row=totals_row, column=1, value="ИТОГО:")
        totals_cell.font = Font(bold=True)
        ws_new.cell(row=totals_row, column=2, value=f"Всего диалогов: {store_rows}")
        ws_new.cell(row=totals_row, column=3, value=f"С галочкой: {store_checkmarks}")
        ws_new.cell(row=totals_row, column=4, value=f"Соотношение: {store_rows} / {store_checkmarks}")
        
        # Set column widths
        ws_new.column_dimensions["A"].width = 12
        ws_new.column_dimensions["B"].width = 25
        ws_new.column_dimensions["C"].width = 20
        ws_new.column_dimensions["D"].width = 30
        
        print(f"  {sheet_name}: {store_rows} диалогов, {store_checkmarks} галочек")
    
    # Save workbook
    wb_new.save(output_file)
    
    print(f"\n{'=' * 80}")
    print(f"Файл сохранён: {output_file.absolute()}")
    print(f"Обработано магазинов: {stores_processed}")
    print(f"Всего записей: {total_rows}")
    print(f"Всего галочек: {total_checkmarks}")
    print(f"\nВнесённые изменения:")
    print(f"  1. Формат времени: ЧЧ:ММ:СС (например, 14:25:07)")
    print(f"  2. Комментарии '1.+ 2.+' заменены на жирную галочку ✔")
    print(f"  3. Итоги добавлены внизу каждого листа магазина")


if __name__ == "__main__":
    create_updated_report()
